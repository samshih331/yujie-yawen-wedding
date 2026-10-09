"""把新人給的 Excel 座位表轉成 seating.js。

用法：~/IdeaProjects/wedding-art/.venv/bin/python tools/build_seating.py "素材/婚禮座位表(1009雲端).xlsx"

表格結構（工作表「婚禮邀請親友」）：
- 第 1–14 列：主桌，名字在 M 欄
- 第 15 列以後：女方在 A–I 欄（桌次 B、名字 D、備註 E），男方在 J–P 欄（桌次 K、名字 M、備註 N）
  每一桌的桌次只寫在該桌第一列，往下到下一個桌次之前都算同一桌
- 同一桌同一個名字出現幾次＝幾個座位（本人＋家人）；名字空白但有填葷/素的座位，算上一位賓客的家人
- 備註寫「某某(素)」表示這個座位坐的是那位賓客（名字欄是同行家人的代表）
只輸出姓名、桌次、人數；桌名、備註、葷素、喜餅都不放上網站。
"""
import json
import re
import sys
from collections import OrderedDict

import openpyxl

ws = openpyxl.load_workbook(sys.argv[1], data_only=True)['婚禮邀請親友']


def clean(v):
    if v is None:
        return ''
    s = re.sub(r'[（(]素[)）]', '', str(v)).strip()
    return re.sub(r'\s+', ' ', s)


seats = OrderedDict()                  # (table, name) -> count

# 新人事後在 LINE 交代、Excel 還沒改的地方：(桌次, Excel 上的名字) -> 正確名字
# （1006 版 Excel 已自行改好 22 桌葛姍姍，目前沒有待補的修正）
FIX = {}
# Excel 上沒有、新人在 LINE 交代要加的賓客：(桌次, 名字, 人數)
EXTRA = [
    # ('17', '楊芝齡', 1),               # 10/8 新娘交代；1008 雲端版 Excel 已自己加上，不再補
]


# 備註裡的稱謂（同學嬤嬤、老公、二舅舅…）不是人名
KIN = r'[嬤媽爸哥姊姐弟妹舅姨伯叔姑嬸婆爺奶孫]|老公|老婆|小孩|嬰兒|同學|朋友|座椅|餐具'


def add(table, name):
    name = FIX.get((table, name), name)
    if name:
        seats[(table, name)] = seats.get((table, name), 0) + 1


for r in range(3, 15):                 # 主桌
    add('主桌', clean(ws.cell(r, 13).value))

for tcol, ncol, ecol, fcol in ((2, 4, 5, 6), (11, 13, 14, 15)):     # 女方、男方（fcol＝葷/素）
    table = last = None
    for r in range(17, ws.max_row + 1):
        t = ws.cell(r, tcol).value
        if isinstance(t, (int, float)):
            table, last = str(int(t)), None
        elif isinstance(t, str) and t.strip() and '不出席' in t:
            break
        if table is None:
            continue
        name = clean(ws.cell(r, ncol).value)
        note = ws.cell(r, ecol).value
        diet = ws.cell(r, fcol).value
        # 備註寫的是坐這個位子的人：寫「某某(素)」，或（1008 版起）備註是 2～4 個字的人名、且這位點素食
        if isinstance(note, str) and (re.search(r'[（(]素[)）]', note)
                                      or (diet == '素' and re.fullmatch(r'[\u4e00-\u9fff]{2,4}', note.strip())
                                          and not re.search(KIN, note))):
            add(table, clean(note))      # 這個座位坐的是備註上的人（名字欄是同行家人的代表）
            last = name or last
        elif name:
            add(table, name)
            last = name
        elif last and ws.cell(r, fcol).value in ('葷', '素'):
            add(table, last)             # 沒寫名字但有點餐的座位，算是上一位賓客的同行家人

for t, n, c in EXTRA:
    seats[(t, n)] = seats.get((t, n), 0) + c

out = [{'name': n, 'table': t, 'count': c} for (t, n), c in seats.items()]
order = lambda g: (-1 if g['table'] == '主桌' else int(g['table']))
out.sort(key=order)

lines = [
    '// 賓客桌次名單，由 tools/build_seating.py 從新人提供的 Excel 座位表產生，請勿手動修改。',
    '// name : 賓客姓名（空白分隔的是別名）；table: 桌次；count: 同名座位數（本人＋家人）',
    'window.SEATING = [',
]
lines += [f'  {json.dumps(g, ensure_ascii=False)},' for g in out]
lines.append('];')
open('seating.js', 'w', encoding='utf-8').write('\n'.join(lines) + '\n')

tables = OrderedDict()
for g in out:
    tables.setdefault(g['table'], []).append(g)
for t, gs in tables.items():
    print(f"{t:>3}  {sum(g['count'] for g in gs):2d} 位  " + '、'.join(f"{g['name']}×{g['count']}" if g['count'] > 1 else g['name'] for g in gs))
print('共', len(out), '筆、', sum(g['count'] for g in out), '個座位、', len(tables), '桌')
